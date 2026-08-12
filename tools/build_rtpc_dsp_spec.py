#!/usr/bin/env python3
"""Build the VehicleSoundSynth V3 RTPC and DSP implementation specification.

Run with a Python environment that provides python-docx:
    python tools/build_rtpc_dsp_spec.py
"""

from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from build_project import ALL_MESSAGES, MACRO_NAMES, STATE_NAMES
from build_user_manual import (
    BLUE,
    CJK_FONT,
    CONTENT_DXA,
    DARK_BLUE,
    INK_BLUE,
    MONO_FONT,
    MUTED,
    TABLE_INDENT_DXA,
    add_bullets,
    add_callout,
    add_code,
    add_heading,
    add_page_field,
    add_steps,
    add_table,
    configure_section,
    configure_styles,
    set_run_font,
)


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "VehicleSoundSynthV3_车机RTPC_DSP链路技术说明.docx"
BASELINE_COMMIT = "81e1345"
SAMPLE_RATE = 48000


def add_paragraph(doc: Document, text: str, *, bold_lead: str | None = None):
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True, color=INK_BLUE)
        rest = paragraph.add_run(text[len(bold_lead) :])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_formula(doc: Document, *lines: str) -> None:
    add_code(doc, "\n".join(lines))


def add_running_furniture(section) -> None:
    section.different_first_page_header_footer = True
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("VehicleSoundSynth V3  |  RTPC 与 DSP 技术说明")
    set_run_font(run, size=8.5, color=MUTED)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    add_page_field(paragraph)
    for run in paragraph.runs:
        set_run_font(run, size=8.5, color=MUTED)


def add_cover(doc: Document) -> None:
    doc.add_paragraph().paragraph_format.space_after = Pt(38)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(6)
    run = kicker.add_run("VEHICLE AUDIO DSP IMPLEMENTATION NOTE")
    set_run_font(run, size=10.5, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(7)
    run = title.add_run("VehicleSoundSynth V3")
    set_run_font(run, size=27, bold=True, color=INK_BLUE)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("车机 RTPC、发声链路与 DSP 数学实现说明")
    set_run_font(run, size=17, bold=True, color=DARK_BLUE)

    add_callout(
        doc,
        "文档口径",
        "只描述当前 V3 已实现事实。基线为分支 codex/vehicle-sound-v3、提交 81e1345，"
        "采样率 48 kHz、Max 8.6.5。V4 规划不在本文范围内。",
    )

    add_table(
        doc,
        ["项目", "值"],
        [
            ["公共声音接口", "vehicle_engine~：1 个消息入口，2 个信号出口（L/R）"],
            ["控制入口", "Internal Manual、Internal Auto、External vss.input"],
            ["音频精度基线", "Max 64-bit 内部计算；运行素材为 48 kHz PCM16 WAV"],
            ["车机公式口径", "Float32/C 参考式 + Q 格式、LUT 与饱和策略"],
            ["版本日期", "2026-08-11"],
        ],
        [2300, 7060],
        font_size=9.4,
    )

    add_paragraph(
        doc,
        "阅读建议：第 1–3 章用于理解 Max patch；第 4–8 章解释声源与混音；"
        "第 9–11 章可直接交给车机 DSP/嵌入式工程师作为移植起点。",
    )
    doc.add_page_break()


def add_architecture(doc: Document) -> None:
    add_heading(doc, "1. 系统边界与完整信号链", 1)
    add_callout(
        doc,
        "一句话结论",
        "这不是单一振荡器直接送入 DAC，而是车辆状态总线驱动多层声源，经过座舱声学、"
        "立体声约束、四段总线 EQ 和峰值保护后输出双声道。",
    )
    add_formula(
        doc,
        "Internal Manual / Internal Auto / External vss.input",
        "                        |",
        "                        v",
        " drive_simulator (internal only) -> vehicle state/event messages",
        "                        |",
        "                        v",
        " engine_controller -> lifecycle + shift envelope + event reset",
        "                        |",
        "        +---------------+---------------+---------------+",
        "        |               |               |               |",
        "        v               v               v               v",
        "  powertrain~      transient~       road_air~     cabin_rumble~",
        " EV/V12/V10+Pad   WAV event family  tire + wind   powertrain excitation",
        "        |               |               |               |",
        "        +---------------+---------------+---------------+",
        "                        | stereo sum",
        "                        v",
        " cabin_master~: macro tone -> 4096-tap IR -> <120 Hz mono -> upper M/S",
        "                        |",
        "                        v",
        " mix_master~: 4-band EQ -> bypass crossfade -> master trim",
        "                        -> omx.peaklim~ -> clip +/-0.891",
        "                        |",
        "                        v",
        " post-EQ Audio gate -> meters/spectrum -> ezdac~ L/R",
    )
    add_table(
        doc,
        ["层级", "实际模块", "输入", "输出/职责"],
        [
            ["驾驶控制", "drive_simulator.maxpat", "油门、刹车、模式、点火、启动、车型", "20 Hz 更新的速度/RPM/挡位及事件消息"],
            ["生命周期", "engine_controller.maxpat", "公共消息", "透传原消息，并生成 starter、engine_running、shift_cut 等内部消息"],
            ["连续动力", "powertrain~.maxpat", "RPM、负载、Sound DNA", "EV/V12/V10、EV Pad、可选参考质感的 L/R"],
            ["瞬态", "transient~.maxpat", "启动、换挡、回火、爆鸣等事件", "真实采样事件的 L/R"],
            ["道路空气", "road_air~.maxpat", "车速、亮度、纹理、宽度", "居中轮胎床 + 独立双声道风噪"],
            ["座舱低频", "cabin_rumble~.maxpat", "Powertrain L/R 与低频宏", "由真实动力信号激励的居中轰鸣"],
            ["座舱与空间", "cabin_master~.maxpat", "stem 总和与宏", "IR、120 Hz 分频、M/S 与早期反射"],
            ["最终总线", "mix_master~.maxpat", "座舱 L/R 与 EQ/Trim", "受保护的最终 L/R"],
        ],
        [1250, 2250, 2450, 3410],
        font_size=8.4,
    )
    add_callout(
        doc,
        "关键依赖",
        "cabin_rumble~ 看起来像并行 stem，但它的激励信号来自 powertrain~ 的 L/R；"
        "它不是独立自由振荡的低频正弦源。",
        kind="caution",
    )


def add_control_flow(doc: Document) -> None:
    add_heading(doc, "2. 控制信号与内部驾驶模型", 1)
    add_heading(doc, "2.1 三种控制来源", 2)
    add_table(
        doc,
        ["来源", "路径", "行为"],
        [
            ["Internal Manual", "UI/W/S -> drive_simulator -> vehicle_engine~", "W：2.5 s 线性加油，松开 1.4 s 回零；S：120 ms 制动，松开立即归零"],
            ["Internal Auto", "20 s 循环驾驶包线 -> drive_simulator", "自动油门、制动、七速 DCT 和内部事件"],
            ["External vss.input", "r vss.input -> switch -> vehicle_engine~", "绕过 drive_simulator；宿主必须发送已经计算好的状态/事件 RTPC"],
        ],
        [1900, 3000, 4460],
        font_size=8.8,
    )
    add_callout(
        doc,
        "外部模式注意",
        "External 模式中 brake、accel、gear 虽被公共接口接受，但当前音频模块没有直接消费它们。"
        "宿主必须同时发送由这些车辆状态推导出的 speed_kph、rpm、load、shift_dir 等有效声学 RTPC。",
        kind="caution",
    )

    add_heading(doc, "2.2 速度积分器", 2)
    add_paragraph(doc, "控制周期 Tctrl = 50 ms，控制频率 20 Hz。所有变量先经过范围裁剪。")
    add_formula(
        doc,
        "u = clamp(throttle, 0, 1)",
        "b = clamp(brake,    0, 1)",
        "v[n+1] = clamp(v[n] + 0.05 * (",
        "           u * (34 - 0.07*v[n])",
        "         - (0.5 + 0.00012*v[n]^2)",
        "         - 55*b), 0, 300)            // speed_kph",
    )
    add_paragraph(
        doc,
        "该模型是听觉演示用离散动力学，不是 SI 单位标定的整车纵向动力学。"
        "空速全油门第一步为 0.05×(34−0.5)=1.675 km/h。",
    )

    add_heading(doc, "2.3 车速到 RPM", 2)
    add_formula(
        doc,
        "rpm_EV  = min(10000, 650 + 31*v + 1400*u)",
        "rpm_V12 = max(780, v * K12[gear])",
        "K12 = {107, 73.5, 55.9, 44.5, 35.6, 29.0, 24.2}",
        "rpm_V10 = max(850, v * K10[gear])",
        "K10 = {114.2, 76.7, 55.5, 43.8, 34.3, 27.7, 22.6}",
    )

    add_heading(doc, "2.4 七速 DCT 与事件", 2)
    add_formula(
        doc,
        "V12 up:   rpm > 6200 + 2600*u and gear < 7",
        "V12 down: rpm < 2200 + 1600*u and gear > 1",
        "V10 up:   rpm > 5800 + 2500*u and gear < 7",
        "V10 down: rpm < 2100 + 1500*u and gear > 1",
        "decision = +1 (up), -1 (down), 0 (hold)",
        "gear_next = clamp(gear + decision, 1, 7)",
    )
    add_bullets(
        doc,
        [
            "换挡判定有 350 ms speedlim 锁定，防止在阈值附近连续抖动。",
            "任意换挡把 shift_cut 设为 0.22，90 ms 后恢复为 1；连续燃烧声短暂下沉。",
            "升挡触发 shift_up 轮询采样，降挡触发 shift_down，并附带 exhaust_pop/crackle。",
            "内部 intake_backfire 在燃油车型、throttle > 0.72、rpm > 2800 的条件上升沿触发，90 ms 后复位。",
            "内部 load = clamp(1.6*u, -1, 1)，accel = clamp(2*u - 2*b, -1, 1)。",
        ],
    )

    add_heading(doc, "2.5 燃油生命周期", 2)
    add_formula(
        doc,
        "drive_enabled = (engine_type == EV) OR (ignition AND start_latched)",
        "start_valid = start AND ignition AND fuel_profile AND start_unlock",
        "run_target = engine_running * combustion_gain * shift_cut",
    )
    add_steps(
        doc,
        [
            "收到有效 start：starter=1、engine_running=0、combustion_gain=0，并锁住重复启动。",
            "500 ms 后：starter=0、engine_running=1、combustion_gain=1。",
            "ignition=0：取消等待任务，立即停止 starter、engine_running 和 combustion_gain，并重新允许下次启动。",
            "降挡或爆鸣请求：exhaust_pop 在 t=0 触发并于 70 ms 复位；crackle 在 120 ms 触发并于 170 ms 复位。",
        ],
    )


def public_rtpc_rows() -> list[list[str]]:
    return [
        ["engine_type", "enum", "ev/v12/v10", "立即", "controller/drive/transient", "选择生命周期与事件车型；实际连续声源混合仍由三个 weight 控制"],
        ["ignition", "bool", "0/1", "立即；Pad 10 ms", "controller/drive/EV Pad", "燃油启动许可；EV Pad gate=5×ignition"],
        ["start", "trigger", "1", "500 ms 状态机", "controller/drive", "燃油 starter -> running 过渡"],
        ["speed_kph", "float", "0..300", "Road 120 ms", "road_air~", "轮胎/风噪幅度；External 模式由宿主提供"],
        ["rpm", "float", "0..10000", "Gen slide；Pad 60 ms", "Gen/EV Pad", "振荡频率、点火频率和 Pad 音高"],
        ["throttle", "float", "0..1", "Gen slide；多处 80–120 ms", "Gen/Pad/reference", "燃烧脉冲锐度、响度、进气、Pad 滤波与质感"],
        ["brake", "float", "0..1", "Internal 50 ms", "drive_simulator", "只通过内部车辆动力学影响声学状态；External 音频无直接消费者"],
        ["load", "float", "-1..1", "Gen 内部", "EV/V12/V10", "EV 取绝对值；燃油仅使用正负载"],
        ["accel", "float", "-1..1", "无", "当前仅契约监视", "当前 V3 没有直接 DSP 影响"],
        ["gear", "int", "1..7", "350 ms DCT 锁定", "drive_simulator", "决定 RPM 比例；送入 engine 后当前无直接 DSP 消费"],
        ["shift", "bool", "0/1", "90 ms cut", "controller/transient", "通用升挡瞬态与燃烧切断"],
        ["shift_dir", "int", "-1/+1", "90 ms cut", "controller/transient", "-1 降挡、+1 升挡；选择不同采样"],
        ["pop", "bool", "0/1", "事件", "controller/transient", "触发 exhaust pop，并安排 crackle"],
        ["intake_backfire", "bool", "0/1", "事件", "transient~", "触发进气回火采样族"],
        ["exhaust_pop", "bool", "0/1", "事件", "controller/transient", "触发排气爆鸣采样并安排 crackle"],
        ["click", "bool", "0/1", "事件", "transient~", "机械换挡机构 click"],
        ["ev_weight", "float", "0..1", "40 ms", "powertrain~", "EV motor+Pad stem 权重；不会自动归一化"],
        ["v12_weight", "float", "0..1", "40 ms", "powertrain/reference", "V12 stem 权重和参考质感门控"],
        ["v10_weight", "float", "0..1", "40 ms", "powertrain/reference", "V10 stem 权重和参考质感门控"],
        ["aggression", "float", "0..1", "Gen/25–60 ms", "Gen/transient/rumble/cabin", "高阶谐波、事件响度、座舱驱动与主级增益"],
        ["brightness", "float", "0..1", "Gen/80 ms", "Gen/road/cabin", "高阶谐波、风噪截止和座舱低通频率"],
        ["body", "float", "0..1", "Gen/60 ms", "Gen/rumble/cabin", "低阶谐波、低频模式和座舱体感"],
        ["texture", "float", "0..1", "Gen/35–80 ms", "Gen/transient/road/cabin", "非整数阶、瞬态、轮胎/风噪和高频细节"],
        ["cabin_width", "float", "0..1", "60–80 ms", "road/cabin", "风噪声道差异和 >120 Hz Side 增益"],
        ["cabin_rumble", "float", "0..1", "60 ms", "cabin_rumble~", "四个座舱低频模态的总深度"],
        ["exhaust_boom", "float", "0..1", "60 ms", "cabin_rumble~", "74 Hz 排气腔体模态"],
        ["structure_vibration", "float", "0..1", "60 ms", "cabin_rumble~", "185 Hz 以下结构/机脚振动"],
        ["cabin_ir_mix", "float", "0..1", "80 ms", "cabin_master~", "等功率 dry=sqrt(1-m)、wet=sqrt(m)"],
        ["powertrain_level_db", "dB", "-60..+6", "50 ms", "powertrain~", "连续动力总 stem 增益"],
        ["transient_level_db", "dB", "-60..+6", "50 ms", "transient~", "所有启动/换挡/爆鸣/机械事件总增益"],
        ["road_level_db", "dB", "-60..+6", "50 ms", "road_air~", "轮胎与风噪总增益"],
        ["rumble_level_db", "dB", "-60..+6", "50 ms", "cabin_rumble~", "座舱轰鸣总增益"],
        ["pad_level_db", "dB", "-60..+6", "50 ms", "ev_pad_driver~", "Pad 独立增益"],
        ["pad_mute", "bool", "0/1", "30 ms", "ev_pad_driver~", "只静音 Pad"],
        ["pad_solo", "bool", "0/1", "30 ms", "powertrain~", "衰减 EV motor，保留 Pad；燃油 stem 不受此控制"],
        ["reference_texture_level_db", "dB", "-60..+6", "50 ms", "reference_texture~", "SF90 residual/fallback 的独立增益"],
        ["reference_texture_mute", "bool", "0/1", "30 ms", "reference_texture~", "静音参考质感"],
        ["eq_low", "list", "30..400 Hz, ±12 dB", "系数更新", "mix_master~", "Low Shelf，默认 90 Hz"],
        ["eq_lowmid", "list", "120..1200 Hz, ±12 dB, Q .3..4", "系数更新", "mix_master~", "Peaking，默认 250 Hz/Q .8"],
        ["eq_highmid", "list", "1..8 kHz, ±12 dB, Q .3..4", "系数更新", "mix_master~", "Peaking，默认 3.8 kHz/Q 1"],
        ["eq_high", "list", "4..16 kHz, ±12 dB", "系数更新", "mix_master~", "High Shelf，默认 9 kHz"],
        ["eq_bypass", "bool", "0/1", "30 ms", "mix_master~", "0=EQ wet，1=线性 crossfade 到 dry"],
        ["master_trim_db", "dB", "-24..0", "50 ms", "mix_master~", "所有 stem、EQ 后的全局输出增益"],
    ]


def add_rtpc_contract(doc: Document) -> None:
    add_heading(doc, "3. RTPC 输入、内部映射与输出关系", 1)
    add_paragraph(
        doc,
        "公共契约由 STATE_NAMES 与 MACRO_NAMES 组成。所有消息进入 vehicle_engine~ 后先经过 "
        "engine_controller；原消息会透传，同时派生内部生命周期和事件消息，再广播给各 stem。",
    )
    add_table(
        doc,
        ["RTPC", "类型", "范围", "平滑/时序", "消费者", "音频映射"],
        public_rtpc_rows(),
        [1300, 720, 1460, 1180, 1840, 2860],
        font_size=7.35,
    )
    add_heading(doc, "3.1 内部控制总线", 2)
    add_table(
        doc,
        ["内部消息", "值/时序", "作用"],
        [
            ["starter", "start 时 1，500 ms 后 0", "选择并播放 V12/V10 starter，off 时停止 play~"],
            ["engine_running", "500 ms 后 1；熄火 0", "门控燃烧 Gen、参考质感和燃油瞬态"],
            ["combustion_gain", "start 时 0；500 ms 后 1", "与 engine_running、shift_cut 相乘形成燃烧包线"],
            ["shift_cut", "正常 1；换挡时 0.22，90 ms 后 1", "短时削减 V12/V10 连续燃烧"],
            ["crackle", "请求后 120 ms 为 1；170 ms 复位", "触发 crackle 轮询采样"],
            ["drive_enabled", "EV=1；燃油=ignition×start", "只门控内部模拟器的有效油门"],
        ],
        [2050, 2700, 4610],
        font_size=8.7,
    )
    add_callout(
        doc,
        "权重与车型不是同一个变量",
        "engine_type 负责生命周期/事件逻辑；ev_weight、v12_weight、v10_weight 负责连续声源混合。"
        "三个权重没有归一化器。量产 RTPC 应维持 one-hot，或由宿主保证总和不超过 1。",
        kind="caution",
    )


def add_powertrain(doc: Document) -> None:
    add_heading(doc, "4. 连续动力合成原理", 1)
    add_heading(doc, "4.1 共同离散算子", 2)
    add_formula(
        doc,
        "clip(x, lo, hi) = min(max(x, lo), hi)",
        "phase[n+1] = frac(phase[n] + f[n]/Fs)",
        "cycle(f)[n] = sin(2*pi*phase[n])",
        "slide: y[n] = y[n-1] + (x[n]-y[n-1])/S",
        "       S = S_up if x[n] > y[n-1] else S_down",
        "dB_to_gain(d) = 10^(d/20) = 2^(d/6.020599913)",
    )
    add_paragraph(
        doc,
        "在 Fs=48 kHz 时，slide 参数 S 以采样数表示；近似 63% 响应时间为 S/Fs，95% 响应时间约为 3S/Fs。",
    )

    add_heading(doc, "4.2 EV Motor Gen", 2)
    add_formula(
        doc,
        "r = slide(clamp(rpm,0,10000), 1200, 1800)",
        "t = slide(clamp(throttle,0,1), 700, 1400)",
        "l = abs(clamp(load,-1,1))",
        "f = clamp(55 + 0.052*r, 55, 760)",
        "a = clamp(0.03 + 0.20*t + 0.05*l, 0, 0.28)",
    )
    add_table(
        doc,
        ["层", "频率", "幅度"],
        [
            ["Coherent low", "f", "0.58 + 0.18×body"],
            ["H2", "2f，110..1520 Hz", "0.08 + 0.05×body"],
            ["H3", "3f，165..2280 Hz", "0.035 + 0.055×texture"],
            ["Inverter", "f×(2.8+brightness)，150..4200 Hz", "0.10 + 0.14×texture"],
            ["Left sheen", "f×(5.2+0.4×brightness)", "0.025×aggression"],
            ["Right sheen", "f×(6.2+0.4×brightness)", "0.025×aggression"],
        ],
        [1900, 3800, 3660],
        font_size=8.7,
    )
    add_formula(
        doc,
        "base = coherent_low + h2 + h3 + inverter",
        "L = tanh(1.2*a*(base + sheen_L)) / 1.2",
        "R = tanh(1.2*a*(coherent_low + .99*h2 + 1.02*h3 + .98*inverter + sheen_R)) / 1.2",
    )
    add_paragraph(
        doc,
        "左右低频基波完全同相，立体差异只放在较高阶 sheen 与轻微幅度差中，从源头避免慢速拍频和低频 AM。",
    )

    add_heading(doc, "4.3 EV 双通道 BEAP Pad", 2)
    add_formula(
        doc,
        "f_pad = clamp(55 + 0.052*rpm, 55, 760)",
        "pitch_CV = log2(f_pad / 261.625565)",
        "osc_A = 2*f_pad                         // +1 octave",
        "osc_B = 3*f_pad                         // +log2(3) octaves",
        "throttle_CV = 0.75 + 3*throttle",
        "filter_CV = (throttle_CV - 0.75)*0.666667 = 2*throttle",
        "gate_CV = 5*ignition",
        "VCA_CV = ADSR * throttle_CV * 0.2",
    )
    add_bullets(
        doc,
        [
            "每个声道包含两只 BEAP Oscillator、Audio Mixer、Ladder、ADSR、VCA。",
            "默认 mixer 为 -15/-18 dB；Ladder 1200 Hz、Res=8；ADSR 450/1000/70/1600 ms。",
            "L/R 的两种 BEAP waveform 枚举次序互换：L=(2,1)，R=(1,2)。枚举语义依赖 Max 8.6.5 BEAP 版本。",
            "120 Hz 以下 Pad 取 (L+R)/2 并复制到双声道；以上使用 82/18 交叉混合。",
        ],
    )
    add_formula(
        doc,
        "PadLow = 0.5*(LP120(L) + LP120(R))",
        "PadHighL = 0.82*HP120(L) + 0.18*HP120(R)",
        "PadHighR = 0.82*HP120(R) + 0.18*HP120(L)",
        "PadOutL/R = dB_to_gain(pad_level_db) * (1-pad_mute) * (PadLow + PadHighL/R)",
    )

    add_heading(doc, "4.4 V12 四冲程双气缸组", 2)
    add_formula(
        doc,
        "r = slide(rpm,700,1100); t = slide(throttle,420,850)",
        "crank = clamp(r/60, 6, 170)",
        "variation = clamp(0.020*sin(2*pi*0.5*crank*t_sec) + 0.008*sin(2*pi*1.5*crank*t_sec), -.03, .03)",
        "bank_fire = clamp(3*crank*(1+variation), 18, 525)",
        "fire = 2*bank_fire ~= 6*rpm/60",
        "phase_R = frac(phase_L + 0.5)",
        "sharpness = 8 + 7*t + 3*aggression",
        "pulse_mean = (1-exp(-sharpness))/sharpness",
        "pulse_L/R = exp(-phase_L/R*sharpness) - pulse_mean",
        "run = slide(engine_running*combustion_gain*shift_cut,480,1900)",
        "a = clamp(.205+.050*body+.30*t+.16*max(load,0),0,.67)*run",
    )
    add_table(
        doc,
        ["阶次", "频率", "幅度"],
        [
            ["0.5", "0.5 fire", ".055+.045 body，随高转衰减"],
            ["1", "fire", ".16+.08 body"],
            ["0.83/1.17", "边带", ".018 + speedmix×(.075/.072)"],
            ["2", "2 fire", ".0945+.108 brightness+.081 speedmix"],
            ["2.33/2.67", "非整数阶", ".045/.040 × speedmix"],
            ["3", "3 fire", ".05625+.06875 texture+.0625 speedmix"],
            ["4/6/8", "高阶", "由 brightness、texture、speedmix 控制"],
            ["1.5", "进气", "throttle×(.025+.040 texture)"],
            ["1.33/2.17", "脉冲门控 residual", "abs(pulse) 门控，texture 控制"],
        ],
        [1500, 2500, 5360],
        font_size=8.4,
    )
    add_formula(
        doc,
        "bank_L = (.88*pulse_L + .12*pulse_R)*cylinder*(.58+.18*body)*(1-.30*speedmix)",
        "bank_R = (.88*pulse_R + .12*pulse_L)*cylinder*(.58+.18*body)*(1-.30*speedmix)",
        "out = clamp(tanh(1.35*drive)/1.25, -.86, .86)",
    )

    add_heading(doc, "4.5 V10 四冲程双气缸组", 2)
    add_formula(
        doc,
        "r = slide(rpm,620,980); t = slide(throttle,360,760)",
        "crank = clamp(r/60, 6, 170)",
        "variation = clamp(.024*sin(2*pi*.5*crank*t_sec)+.006*sin(2*pi*2.5*crank*t_sec),-.03,.03)",
        "bank_fire = clamp(2.5*crank*(1+variation),16,450)",
        "fire = 2*bank_fire ~= 5*rpm/60",
        "sharpness = 7 + 9*t + 4*aggression",
        "run = slide(engine_running*combustion_gain*shift_cut,420,1650)",
        "a = clamp(.225+.055*body+.32*t+.18*max(load,0),0,.70)*run",
    )
    add_table(
        doc,
        ["阶次", "频率", "幅度"],
        [
            ["0.5", "0.5 fire", ".060+.052 body，随高转衰减"],
            ["1", "fire", ".17+.08 aggression"],
            ["0.81/1.19", "边带", ".022 + speedmix×(.080/.078)"],
            ["1.5", "奇数阶进气色彩", ".040+.050 texture"],
            ["2", "2 fire", ".08775+.0945 brightness+.07425 speedmix"],
            ["2.2/2.8", "非整数阶", ".070/.060 × speedmix"],
            ["2.5", "V10 rasp", ".055+.075 texture+.060 speedmix"],
            ["3/5/7", "高阶", "aggression/brightness/texture 与 speedmix"],
            ["1.25 FM", "fire×1.25×(1+.008 sin(crank))", "throttle×(.035+.055 texture)"],
            ["1.37/2.23", "脉冲门控 residual", "abs(pulse) 门控，texture 控制"],
        ],
        [1500, 3000, 4860],
        font_size=8.2,
    )
    add_formula(
        doc,
        "bank_L = (.86*pulse_L + .14*pulse_R)*cylinder*(.60+.18*body)*(1-.32*speedmix)",
        "bank_R = (.86*pulse_R + .14*pulse_L)*cylinder*(.60+.18*body)*(1-.32*speedmix)",
        "out = clamp(tanh(1.42*drive)/1.27, -.88, .88)",
    )

    add_heading(doc, "4.6 可选 Ferrari Reference Texture", 2)
    add_formula(
        doc,
        "profile_gate = engine_running * clamp(v12_weight + v10_weight, 0, 1)",
        "throttle_lift = 10^((4*throttle)/20)       // 0..+4 dB",
        "texture_gain = profile_gate * throttle_lift * dB_to_gain(reference_texture_level_db)",
        "output = groove_loop * texture_gain * (1-reference_texture_mute)",
    )
    add_paragraph(
        doc,
        "优先加载 media/local/ferrari_sf90_residual.wav；不存在时加载确定性 fallback。"
        "本地残差为 250 ms 双声道循环、-18 dBFS RMS、-3 dBFS 峰值上限，不得提交或分发。",
    )


def add_transients_and_noise(doc: Document) -> None:
    add_heading(doc, "5. 瞬态采样、道路与空气层", 1)
    add_heading(doc, "5.1 瞬态采样", 2)
    add_table(
        doc,
        ["事件族", "变体", "播放前固定增益", "素材峰值归一化"],
        [
            ["V12/V10 starter", "各 1", "0.26", "-3 dBFS"],
            ["shift_up", "3", "0.32", "-6 dBFS"],
            ["shift_down", "3", "0.38", "-6 dBFS"],
            ["intake_backfire", "3", "0.34", "-6 dBFS"],
            ["exhaust_pop", "4", "0.30", "-5 dBFS"],
            ["exhaust_crackle", "3", "0.18", "-8 dBFS"],
            ["mechanical_click", "1", "0.18", "-6 dBFS"],
        ],
        [2400, 1100, 2700, 3160],
        font_size=8.8,
    )
    add_formula(
        doc,
        "event_gain = (0.72 + 0.38*aggression) * (0.72 + 0.46*texture)",
        "L = event_sum * event_gain * 0.94 * dB_to_gain(transient_level_db)",
        "R = event_sum * event_gain * 0.88 * dB_to_gain(transient_level_db)",
    )
    add_paragraph(
        doc,
        "素材来自六个已记录来源中的 CC0 录音，处理为 48 kHz PCM16、去 DC、短淡入淡出并按峰值归一化。"
        "连续发动机仍是程序化合成，不能把瞬态素材误称为多负载循环发动机样本库。",
    )

    add_heading(doc, "5.2 Road/Air 白噪声来源", 2)
    add_formula(
        doc,
        "v = clamp(speed_kph, 0, 300)",
        "A_tire = 0.065 * (v/300)^0.8",
        "A_wind = 0.045 * (v/300)^1.7",
        "f_wind_LP = 1400 + 2600*brightness",
        "G_tire_texture = 0.58 + 0.32*texture",
        "G_wind_texture = 0.30 + 0.48*texture",
        "G_wind_width = 0.55 + 0.45*cabin_width",
    )
    add_bullets(
        doc,
        [
            "轮胎层使用一个居中的 pink~，带宽约 70–900 Hz。",
            "风噪使用两个相互独立的 noise~，高通 400 Hz、低通约 1.4–4 kHz。",
            "V12/V10 Gen 内部没有 noise()；听到的持续白噪感主要来自 road_air~ 风噪。",
            "最终 Road/Air stem 乘以 dB_to_gain(road_level_db)，默认 V12/V10 为 -16 dB。",
        ],
    )


def add_cabin(doc: Document) -> None:
    add_heading(doc, "6. 座舱低频轰鸣与结构振动", 1)
    add_formula(
        doc,
        "mid = 0.5*(powertrain_L + powertrain_R)",
        "excitation = LP520(HP28(mid))",
        "G_cabin    = 0.82*cabin_rumble",
        "G_exhaust  = 0.62*exhaust_boom",
        "G_structure= 0.58*structure_vibration",
        "G_body     = 0.72 + 0.48*body",
        "G_drive    = 0.82 + 0.28*aggression",
    )
    add_table(
        doc,
        ["分量", "滤波/模态", "固定权重", "宏控制"],
        [
            ["Cabin mode 1", "48 Hz SVF bandpass", "0.14", "cabin_rumble"],
            ["Cabin mode 2", "68 Hz SVF bandpass", "0.42", "cabin_rumble"],
            ["Cabin mode 3", "92 Hz SVF bandpass", "0.34", "cabin_rumble"],
            ["Cabin mode 4", "118 Hz SVF bandpass", "0.22", "cabin_rumble"],
            ["Structure", "onepole~ 185", "—", "structure_vibration"],
            ["Exhaust", "74 Hz SVF lowpass output", "—", "exhaust_boom"],
        ],
        [2200, 3100, 1500, 2560],
        font_size=8.8,
    )
    add_formula(
        doc,
        "rumble = clip((structure + cabin_modes + exhaust) * G_body * G_drive, -.72, .72)",
        "L = R = rumble * dB_to_gain(rumble_level_db)",
    )
    add_callout(
        doc,
        "听觉含义",
        "该 stem 生成车内可感知的机脚、车体和排气腔体能量，最终以完全居中的低频送入座舱总线。"
        "它不会制造与发动机无关的自由运行正弦。",
    )

    add_heading(doc, "7. 座舱 IR 与双声道空间处理", 1)
    add_heading(doc, "7.1 宏映射", 2)
    add_formula(
        doc,
        "G_aggression = 0.48 + 0.20*aggression",
        "f_brightness = 4500 + 14500*brightness",
        "G_body       = 0.78 + 0.25*body",
        "G_texture    = 0.94 + 0.06*texture",
        "W_side       = 0.60 + 1.20*cabin_width",
    )
    add_heading(doc, "7.2 4096-tap 座舱 IR", 2)
    add_formula(
        doc,
        "g_dry = sqrt(1-m); g_wet = sqrt(m); m = clamp(cabin_ir_mix,0,1)",
        "wet_L[n] = sum(k=0..4095, h_L[k]*x_L[n-k])",
        "wet_R[n] = sum(k=0..4095, h_R[k]*x_R[n-k])",
        "mix_L/R = g_dry*x_L/R + g_wet*wet_L/R",
    )
    add_heading(doc, "7.3 120 Hz 以下居中、以上 M/S", 2)
    add_formula(
        doc,
        "Low = 0.5*(LP120(mix_L) + LP120(mix_R))",
        "HL = allpass(delay=1.7 ms, feedback=.42, HP120(mix_L))",
        "HR = allpass(delay=2.3 ms, feedback=-.36, HP120(mix_R))",
        "Mid  = 0.5*(HL+HR)",
        "Side = 0.5*(HL-HR)",
        "Side_safe = 0.85*W_side*Side",
        "L = Low + Mid + Side_safe + 0.12*delay(HR,0.65 ms)",
        "R = Low + Mid - Side_safe + 0.12*delay(HL,1.35 ms)",
    )
    add_paragraph(
        doc,
        "因此低于 120 Hz 的深低频严格相同，立体相位差和早期反射只进入高频分支。"
        "固定 0.85 Side 安全系数用于降低折叠到 mono 时的损失。",
    )


def add_master_eq(doc: Document) -> None:
    add_heading(doc, "8. 全局 EQ、总线增益与输出", 1)
    add_table(
        doc,
        ["节点", "类型", "默认", "可调范围"],
        [
            ["0", "Low Shelf", "90 Hz / 0 dB / S=.7", "30–400 Hz，±12 dB"],
            ["1", "Peaking", "250 Hz / 0 dB / Q=.8", "120–1200 Hz，±12 dB，Q .3–4"],
            ["2", "Peaking", "3800 Hz / 0 dB / Q=1", "1–8 kHz，±12 dB，Q .3–4"],
            ["3", "High Shelf", "9000 Hz / 0 dB / S=.7", "4–16 kHz，±12 dB"],
        ],
        [900, 1800, 3100, 3560],
        font_size=8.8,
    )
    add_formula(
        doc,
        "EQ wet: cascade of 4 biquad sections, identical coefficients for L/R",
        "bypass dry_gain = eq_bypass",
        "bypass wet_gain = 1 - eq_bypass              // 30 ms line~",
        "trim_gain = 10^(master_trim_db/20)            // -24..0 dB, 50 ms",
        "master = trim_gain*(dry_gain*x + wet_gain*EQ(x))",
        "protected = clip(omx_peaklim(master), -0.891, +0.891)",
        "monitor = protected * linear_ramp(Audio, 60 ms)",
    )
    add_callout(
        doc,
        "不可误解",
        "filtergraph~ 只计算/显示滤波器系数；真正处理音频的是左右两个 cascade~ 4。"
        "Master Trim 位于 EQ/bypass 之后，影响所有 stem。",
    )


def add_chip_math(doc: Document) -> None:
    add_heading(doc, "9. 车机芯片可执行数学形式", 1)
    add_heading(doc, "9.1 推荐数值格式", 2)
    add_table(
        doc,
        ["数据", "Float32 参考", "定点建议", "保护"],
        [
            ["throttle/brake/macros", "float [0,1]", "Q1.31", "入口 clamp"],
            ["load/accel/audio sample", "float [-1,1]", "Q1.31", "每个非线性/输出级饱和"],
            ["speed/RPM/frequency", "float", "Q16.16", "上限 300/10000/Fs/2"],
            ["oscillator phase", "float [0,1)", "Q0.32", "无符号自然回绕"],
            ["gain/biquad coefficient", "float", "Q2.30", "64-bit MAC；逐节缩放"],
            ["mix accumulator", "float", "64-bit Q5.59 或更宽", "总线后统一舍入/饱和"],
        ],
        [2200, 1800, 2300, 3060],
        font_size=8.6,
    )
    add_formula(
        doc,
        "phase_inc_Q32 = round(f_Hz * 2^32 / Fs)",
        "phase_Q32 = phase_Q32 + phase_inc_Q32          // unsigned wrap",
        "sin_value_Q31 = sin_LUT(phase_Q32)",
    )

    add_heading(doc, "9.2 指数、对数和频谱坐标", 2)
    add_formula(
        doc,
        "dB_to_gain(d) = 10^(d/20) = 2^(d/6.020599913)",
        "gain_to_dB(g) = 20*log10(max(g,1e-6))",
        "pitch_octaves = log2(f/261.625565)",
        "log_axis(f) = log10(f/f_min) / log10(f_max/f_min)",
        "f_from_axis(x) = f_min * (f_max/f_min)^x",
        "2*10^1 = 20 Hz; 2*10^2 = 200 Hz; 2*10^3 = 2 kHz; 2*10^4 = 20 kHz",
    )
    add_bullets(
        doc,
        [
            "定点芯片可把 -60..+6 dB 以 0.25 dB 步长预计算为 exp2 LUT，再对相邻表项线性插值。",
            "log2 可先用 count-leading-zeros 提取指数，再对 mantissa 查 256 点 LUT。",
            "exp(-k×phase) 可改写为 2^(-k×phase/log(2))，复用 exp2 LUT。",
            "sqrt 用 256 点 LUT 或整数牛顿迭代；tanh 用对称 LUT，输入先限制在 [-4,4]。",
        ],
    )

    add_heading(doc, "9.3 四段 EQ 的标准 biquad 形式", 2)
    add_paragraph(
        doc,
        "Max 当前由 filtergraph~ 生成系数。车机参考采用 RBJ/W3C Audio EQ Cookbook 的等效 biquad；"
        "两者应通过脉冲响应和扫频对齐，不能假定系数位级一致。",
    )
    add_formula(
        doc,
        "w0 = 2*pi*f0/Fs; c = cos(w0); s = sin(w0)",
        "A = 10^(dBgain/40)",
        "alpha_peak = s/(2*Q)",
        "alpha_shelf = s/2 * sqrt((A+1/A)*(1/S-1)+2)",
        "beta = 2*sqrt(A)*alpha_shelf",
        "y[n] = B0*x[n] + B1*x[n-1] + B2*x[n-2] - A1*y[n-1] - A2*y[n-2]",
        "where Bk=bk/a0 and Ak=ak/a0",
    )
    add_heading(doc, "Peaking", 3)
    add_formula(
        doc,
        "b0=1+alpha*A; b1=-2*c; b2=1-alpha*A",
        "a0=1+alpha/A; a1=-2*c; a2=1-alpha/A",
    )
    add_heading(doc, "Low Shelf", 3)
    add_formula(
        doc,
        "b0=A*((A+1)-(A-1)*c+beta)",
        "b1=2*A*((A-1)-(A+1)*c)",
        "b2=A*((A+1)-(A-1)*c-beta)",
        "a0=(A+1)+(A-1)*c+beta",
        "a1=-2*((A-1)+(A+1)*c)",
        "a2=(A+1)+(A-1)*c-beta",
    )
    add_heading(doc, "High Shelf", 3)
    add_formula(
        doc,
        "b0=A*((A+1)+(A-1)*c+beta)",
        "b1=-2*A*((A-1)+(A+1)*c)",
        "b2=A*((A+1)+(A-1)*c-beta)",
        "a0=(A+1)-(A-1)*c+beta",
        "a1=2*((A-1)-(A+1)*c)",
        "a2=(A+1)-(A-1)*c-beta",
    )
    add_callout(
        doc,
        "定点 biquad",
        "系数建议 Q2.30、状态和乘加使用 64 位；每一节单独归一化 a0，四节串联。"
        "对 ±12 dB 极值、16 kHz 高架和 Q=4 做溢出及极点稳定性测试。",
        kind="caution",
    )


def add_portability(doc: Document) -> None:
    add_heading(doc, "10. Max 到车机的可移植性边界", 1)
    add_table(
        doc,
        ["等级", "模块/对象", "处理策略"],
        [
            ["可按公式直接移植", "drive_simulator、engine_controller、Gen 振荡器、dB、M/S、FIR", "实现公式并用 golden vector 验证"],
            ["等效实现后对齐", "svf~、lores~、onepole~、cross~、allpass~", "选定车机 IIR/FIR 结构，扫频对齐幅相响应"],
            ["需要冻结或替换", "BEAP Oscillator/Ladder/ADSR/VCA", "冻结用户调音后导出参数；用目标芯片模块复现并做 A/B"],
            ["Max 黑盒替换", "omx.peaklim~", "选择目标平台 limiter；重新标定阈值、attack/release 和总延迟"],
            ["直接携带资源", "事件 WAV、cabin IR、可分发 fallback", "按 manifest 校验 SHA-256、采样率、声道和许可证"],
            ["禁止发布", "Ferrari.wav 与 media/local residual", "只可本机参考，不能提交、同步或进入量产包"],
        ],
        [1700, 3000, 4660],
        font_size=8.55,
    )
    add_callout(
        doc,
        "真实性边界",
        "当前系统的启动、换挡、回火、爆鸣和座舱 IR 使用实录素材；"
        "V12/V10 连续运转是程序化声学原型，不是整车多转速、多负载、多麦克风采样库。",
        kind="risk",
    )


def add_golden_vectors(doc: Document) -> None:
    add_heading(doc, "11. 数值样例与验收向量", 1)
    db_values = [-60, -24, -8, -3, 0, 6]
    add_table(
        doc,
        ["dB", "线性增益 10^(dB/20)"],
        [[f"{value:+d}", f"{10 ** (value / 20):.9f}"] for value in db_values],
        [2600, 6760],
        font_size=9.4,
    )
    add_table(
        doc,
        ["场景", "计算", "期望值"],
        [
            ["EV 650 rpm", "55+0.052×650", "88.8 Hz"],
            ["EV 3000 rpm", "55+0.052×3000", "211 Hz"],
            ["EV 10000 rpm", "55+0.052×10000", "575 Hz"],
            ["V12 idle 780 rpm", "780/60×6", "78 Hz 主点火事件频率"],
            ["V12 3500 rpm", "3500/60×6", "350 Hz"],
            ["V10 idle 850 rpm", "850/60×5", "70.8333 Hz"],
            ["V10 4200 rpm", "4200/60×5", "350 Hz"],
            ["空速全油门首个 tick", ".05×(34-.5)", "+1.675 km/h"],
            ["Pad exact fifth", "oscA=2f，oscB=3f", "频率比 3:2"],
            ["输出硬夹幅", "20log10(.891)", f"{20*math.log10(.891):.3f} dBFS"],
        ],
        [3100, 3100, 3160],
        font_size=8.6,
    )

    presets = json.loads((ROOT / "presets" / "sound_dna.json").read_text(encoding="utf-8"))
    add_heading(doc, "11.1 三个默认 Sound DNA", 2)
    add_table(
        doc,
        ["参数", "Future EV", "High-Rev V12", "Aggressive V10"],
        [
            ["engine_type", presets["Future EV"]["engine_type"], presets["High-Rev V12"]["engine_type"], presets["Aggressive V10"]["engine_type"]],
            ["aggression", str(presets["Future EV"]["aggression"]), str(presets["High-Rev V12"]["aggression"]), str(presets["Aggressive V10"]["aggression"])],
            ["brightness", str(presets["Future EV"]["brightness"]), str(presets["High-Rev V12"]["brightness"]), str(presets["Aggressive V10"]["brightness"])],
            ["body", str(presets["Future EV"]["body"]), str(presets["High-Rev V12"]["body"]), str(presets["Aggressive V10"]["body"])],
            ["texture", str(presets["Future EV"]["texture"]), str(presets["High-Rev V12"]["texture"]), str(presets["Aggressive V10"]["texture"])],
            ["cabin_rumble", str(presets["Future EV"]["cabin_rumble"]), str(presets["High-Rev V12"]["cabin_rumble"]), str(presets["Aggressive V10"]["cabin_rumble"])],
            ["powertrain dB", str(presets["Future EV"]["powertrain_level_db"]), str(presets["High-Rev V12"]["powertrain_level_db"]), str(presets["Aggressive V10"]["powertrain_level_db"])],
            ["transient dB", str(presets["Future EV"]["transient_level_db"]), str(presets["High-Rev V12"]["transient_level_db"]), str(presets["Aggressive V10"]["transient_level_db"])],
            ["road/air dB", str(presets["Future EV"]["road_level_db"]), str(presets["High-Rev V12"]["road_level_db"]), str(presets["Aggressive V10"]["road_level_db"])],
            ["rumble dB", str(presets["Future EV"]["rumble_level_db"]), str(presets["High-Rev V12"]["rumble_level_db"]), str(presets["Aggressive V10"]["rumble_level_db"])],
            ["master trim dB", str(presets["Future EV"]["master_trim_db"]), str(presets["High-Rev V12"]["master_trim_db"]), str(presets["Aggressive V10"]["master_trim_db"])],
        ],
        [2100, 2420, 2420, 2420],
        font_size=8.55,
    )


def add_sources(doc: Document) -> None:
    add_heading(doc, "12. 事实来源与实现入口", 1)
    add_table(
        doc,
        ["类别", "来源"],
        [
            ["项目生成事实", "tools/build_project.py；生成后的 patchers/*.maxpat 与 gen/*.gendsp"],
            ["公共使用契约", "README.md、presets/sound_dna.json、tests/test_project.py"],
            ["素材来源与归一化", "assets/asset_manifest.json、ASSET_LICENSES.md、tools/process_assets.py"],
            ["Gen slide 定义", "https://docs.cycling74.com/reference/gen_dsp_slide/"],
            ["filtergraph~/cascade~", "https://docs.cycling74.com/max8/refpages/filtergraph~"],
            ["芯片 EQ 参考", "https://www.w3.org/TR/audio-eq-cookbook/"],
        ],
        [2300, 7060],
        font_size=8.8,
    )
    add_callout(
        doc,
        "工程交付建议",
        "移植时先实现控制状态机和 EV/V12/V10 dry stem，再逐层加入瞬态、rumble、IR、M/S 和 EQ。"
        "每增加一层都保存双声道 48 kHz golden render，比较 RMS、频谱、相关性、mono fold 和峰值。",
    )


def audit_document(path: Path) -> None:
    with ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        styles_xml = archive.read("word/styles.xml").decode("utf-8")
        numbering_xml = archive.read("word/numbering.xml").decode("utf-8")
        settings_xml = archive.read("word/settings.xml").decode("utf-8")

    missing_messages = [name for name in ALL_MESSAGES if name not in document_xml]
    checks = {
        "Letter page width": 'w:w="12240"' in document_xml,
        "Letter page height": 'w:h="15840"' in document_xml,
        "one-inch margins": all(token in document_xml for token in ('w:top="1440"', 'w:right="1440"', 'w:bottom="1440"', 'w:left="1440"')),
        "compact body line spacing": 'w:line="300"' in styles_xml,
        "heading color": BLUE in styles_xml,
        "CJK font override": CJK_FONT in styles_xml,
        "real numbering": "<w:abstractNum" in numbering_xml and "<w:numPr>" in document_xml,
        "fixed table width": f'w:w="{CONTENT_DXA}"' in document_xml,
        "table indent": f'w:w="{TABLE_INDENT_DXA}"' in document_xml,
        "field update": "w:updateFields" in settings_xml,
        "all public messages": not missing_messages,
        "core equations": all(token in document_xml for token in ("dB_to_gain", "phase_inc_Q32", "v[n+1]", "bank_fire", "Side_safe")),
        "no placeholders": not re.search(r"\b(TODO|TBD|PLACEHOLDER)\b", document_xml),
    }
    failures = [name for name, ok in checks.items() if not ok]
    if missing_messages:
        failures.append("missing messages: " + ", ".join(missing_messages))
    if failures:
        raise RuntimeError("DOCX audit failed: " + "; ".join(failures))


def build_document(output: Path) -> None:
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    add_running_furniture(doc.sections[0])

    add_cover(doc)
    add_architecture(doc)
    add_control_flow(doc)
    add_rtpc_contract(doc)
    add_powertrain(doc)
    add_transients_and_noise(doc)
    add_cabin(doc)
    add_master_eq(doc)
    add_chip_math(doc)
    add_portability(doc)
    add_golden_vectors(doc)
    add_sources(doc)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name is None:
                            set_run_font(run)

    core = doc.core_properties
    core.title = "VehicleSoundSynth V3 车机 RTPC、发声链路与 DSP 数学实现说明"
    core.subject = "当前 V3 的控制总线、EV/V12/V10 合成、瞬态、座舱、EQ 与车机定点移植公式"
    core.author = "VehicleSoundSynth Project"
    core.keywords = "Max/MSP, Vehicle Sound, RTPC, DSP, V12, V10, EV, Fixed Point, Biquad"
    core.comments = f"Generated from VehicleSoundSynth V3 baseline {BASELINE_COMMIT}; V4 is out of scope."

    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    audit_document(output)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    output = args.output.resolve()
    build_document(output)
    print(output)


if __name__ == "__main__":
    main()
